"""
Lecture des fichiers de BOM (XLSX EasyEDA, CSV KiCad/Altium…) et heuristiques
de détection des colonnes.

Deux pièges réels traités ici — ce ne sont pas des raffinements, ce sont les
deux façons dont une BOM ment silencieusement :

  A. la quantité fiable est le NOMBRE DE DESIGNATORS, pas la colonne
     Quantity (souvent fausse après édition manuelle du fichier) ;
  B. les lignes DNP (do not populate) non détectées produisent un écart de
     quantité invisible au moment d'acheter.
"""
import csv
import io
import re

from openpyxl import load_workbook

# Champs canoniques proposés sur l'écran de mapping.
CANONICAL_FIELDS = [
    ("designator", "Designators (R1, C3…)"),
    ("value", "Valeur"),
    ("package", "Boîtier / footprint"),
    ("mpn", "Réf fabricant (MPN)"),
    ("manufacturer", "Fabricant"),
    ("lcsc", "Réf LCSC (C…)"),
    ("qty", "Quantité"),
    ("dnp", "DNP / non monté"),
    ("ignore", "— ignorer cette colonne —"),
]

# Entête → champ canonique. L'ORDRE EST LA PRIORITÉ (1er motif gagnant) :
# `mpn` doit passer AVANT `manufacturer`, sinon la colonne « Manufacturer
# Part Number » serait classée comme fabricant.
HEADER_PATTERNS = [
    ("lcsc", r"lcsc|supplier\s*part|jlc"),
    ("mpn", r"manufacturer\s*part|mfr.*part|mpn|part\s*number|p/n"),
    ("manufacturer", r"manufacturer|mfr|fabricant|brand"),
    ("designator", r"designator|reference|ref\b|refdes"),
    ("package", r"package|footprint|boitier|boîtier|case"),
    ("qty", r"qty|quantity|quantit"),
    ("dnp", r"dnp|no[tn].*(fit|populate)|bom\s*excl"),
    ("value", r"value|valeur|comment"),
]

# Est MONTÉ seulement si la cellule DNP vaut l'une de ces valeurs ; tout le
# reste est considéré comme non monté (on préfère alerter à tort qu'acheter
# trop peu).
FITTED_VALUES = ("", "0", "false", "no", "non", "fitted", "fit")


def read_table(uploaded_file):
    """Fichier XLSX ou CSV → (entêtes, lignes de dicts {entête: valeur str}).
    Tout est converti en CHAÎNES : la zone de transit stocke le brut,
    l'interprétation vient après le mapping."""
    if uploaded_file.name.lower().endswith((".xlsx", ".xlsm")):
        ws = load_workbook(uploaded_file, read_only=True, data_only=True).active
        rows = [[("" if c is None else str(c).strip()) for c in r]
                for r in ws.iter_rows(values_only=True)]
    else:
        raw = uploaded_file.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        # Séparateur détecté automatiquement : virgule, point-virgule, tab.
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
        rows = [[c.strip() for c in r]
                for r in csv.reader(io.StringIO(text), dialect)]
    rows = [r for r in rows if any(r)]          # lignes vides éliminées
    if not rows:
        return [], []
    headers = [h or f"colonne {i + 1}" for i, h in enumerate(rows[0])]
    return headers, [dict(zip(headers, r)) for r in rows[1:]]


def guess_mapping(headers):
    """Propose un mapping entête → champ canonique (l'utilisateur corrige)."""
    mapping = {}
    for h in headers:
        mapping[h] = next((field for field, pattern in HEADER_PATTERNS
                           if re.search(pattern, h, re.IGNORECASE)), "ignore")
    return mapping


def extract_lines(raw_rows, mapping):
    """
    Applique le mapping et produit les champs canoniques de chaque ligne,
    avec un éventuel `warning` de divergence quantité/designators.
    """
    # Champ canonique → nom de colonne source (la première trouvée).
    col_for = {}
    for header, field in mapping.items():
        col_for.setdefault(field, header)

    out = []
    for idx, row in enumerate(raw_rows):
        def get(field):
            return row.get(col_for.get(field, ""), "").strip()

        designators = get("designator")
        # « R1,R2,R7 » ne ment pas : le nombre de designators fait foi.
        n_des = len([d for d in re.split(r"[,;\s]+", designators) if d])
        try:
            qty_col = float(get("qty").replace(",", ".") or 0)
        except ValueError:
            qty_col = 0
        warning = ""
        if n_des and qty_col and n_des != qty_col:
            warning = (f"quantité ({qty_col:g}) ≠ nb de designators "
                       f"({n_des}) — on retient {n_des}")
        out.append({
            "row_index": idx,
            "designators": designators[:300],
            "qty": n_des or qty_col or 1,
            "value": get("value")[:80],
            "package": get("package")[:40],
            "mpn": get("mpn")[:120],
            "manufacturer": get("manufacturer")[:120],
            "lcsc": get("lcsc")[:40],
            "dnp": get("dnp").lower() not in FITTED_VALUES,
            "warning": warning,
        })
    return out
